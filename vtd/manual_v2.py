"""MANUAL v2: conservative semantic contracts and fail-closed rendering.

This module is intentionally independent of model/video execution.  The legacy
pipeline can adapt its extracted signals here before rendering.
"""
from dataclasses import asdict
from html import escape
from pathlib import Path
from typing import Any, Iterable, List, Optional

from .action_qualification import (
    CONFIRMED_ACTION, EXPLICIT_CONFIGURATION_INSTRUCTION, classify_candidate,
)
from .contracts import Evidence, ManualDocument, ManualStepContract, ProcedureGroup, SemanticCandidate
from .evidence_selector import select_evidence_frame


def _ui_target(text: str) -> str:
    """Return the concrete object named after the imperative verb."""
    import re
    match = re.search(r"\b(?:kliknij|naciśnij|wybierz|otwórz|zamknij|wpisz|wprowadź|zaznacz|odznacz|przejdź|uruchom|zapisz|usuń|dodaj|ustaw|zmień|wyślij|zatwierdź|zatwierdzamy|zatwierdzimy)\b(.+)", text, re.I)
    if not match:
        return "Konfiguracja — nazwa pola/wartości do potwierdzenia"
    target = match.group(1).strip(" .,:;—-()")
    return target[:120]


def _text(item: Any) -> str:
    """Use enriched operation wording, retaining speech as fallback/context."""
    if isinstance(item, dict):
        enriched = " ".join(str(item.get(k) or "").strip() for k in ("title", "description") if item.get(k))
        return enriched or str(item.get("text", item.get("speech_text", "")) or "").strip()
    for key in ("enriched_title", "enriched_description", "text", "speech_text"):
        value = getattr(item, key, None)
        if value:
            return str(value).strip()
    return ""


def _speech_context(item: Any) -> str:
    if isinstance(item, dict):
        return str(item.get("speech_text", item.get("text", "")) or "").strip()
    return str(getattr(item, "speech_text", "") or "").strip()


def _time(item: Any, name: str, default: float = 0.0) -> float:
    value = getattr(item, name, None)
    if value is None and name == "start":
        value = getattr(item, "start_time", None)
    if value is None and name == "end":
        value = getattr(item, "end_time", None)
    if value is None and isinstance(item, dict):
        value = item.get(name)
        if value is None and name == "start":
            value = item.get("start_time")
        if value is None and name == "end":
            value = item.get("end_time")
    return float(value if value is not None else default)


def _candidate(item: Any) -> SemanticCandidate:
    if isinstance(item, SemanticCandidate):
        return item
    text = _text(item)
    get = (lambda k, d=None: getattr(item, k, d)) if not isinstance(item, dict) else (lambda k, d=None: item.get(k, d))
    title = get("title", get("enriched_title"))
    description = get("description", get("enriched_description"))
    if title or description:
        # Description carries the operational sentence; title stays explicit metadata.
        text = str(description or title).strip()
    local_context = str(get("local_context", "") or "")
    speech = _speech_context(item)
    if speech and speech != text:
        local_context = (local_context + " " + speech).strip()
    return SemanticCandidate(text=text, title=title, description=description, source_segment_ids=list(get("source_segment_ids", []) or []), frame_ids=list(get("frame_ids", []) or []), curator_decision=str(get("curator_decision", "unknown")), enriched_step_id=get("enriched_step_id"), qualification=str(get("qualification", "unknown")), evidence_status=str(get("evidence_status", "missing")), semantic_intent=get("semantic_intent"), procedure_id=get("procedure_id"), start=_time(item, "start"), end=_time(item, "end", _time(item, "start") + 1), local_context=local_context, visible_ui=get("visible_ui"))


def _dedup_key(c: SemanticCandidate) -> tuple:
    import re
    words = re.findall(r"\w+", c.text.lower())
    return (c.semantic_intent or c.procedure_id or "", tuple(words[:20]), tuple(c.source_segment_ids))


def build_manual_v2(title: str, candidates: Iterable[Any], *, goal: str = "DO POTWIERDZENIA",
                    preconditions: Optional[List[str]] = None,
                    evidence_by_candidate: Optional[Iterable[Iterable[dict]]] = None) -> ManualDocument:
    """Convert transcript/action candidates into grouped, qualified manual steps.

    Commentary, scene changes and unsupported model suggestions are excluded.
    Evidence is accepted only when the selector confirms QA and semantic/OCR
    alignment; no midpoint fallback is used.
    """
    raw_candidates = [_candidate(item) for item in candidates]
    candidates = []
    seen = set()
    for candidate in raw_candidates:
        key = _dedup_key(candidate)
        if key in seen and key[0]:
            continue
        seen.add(key); candidates.append(candidate)
    evidence_rows = list(evidence_by_candidate or [])
    steps: List[ManualStepContract] = []
    rejected_candidates = []
    for index, candidate in enumerate(candidates):
        text = candidate.text
        # Qualification happens after curation/enrichment, over the local window.
        context = candidate.local_context or text
        result = classify_candidate(text, context, candidate.visible_ui)
        start, end = candidate.start, candidate.end
        if end < start: end = start
        # Fusion's ManualStep is also the authoritative logical-step stream.  A
        # step with no speech is not silently discarded: retain it as REJECTED
        # so the renderer cannot claim an empty procedure or VALID content.
        if result["classification"] not in (CONFIRMED_ACTION, EXPLICIT_CONFIGURATION_INSTRUCTION):
            rejected_text = text or "Materiał bez jawnej czynności — odrzucony do procedury"
            rejected = ManualStepContract(action=rejected_text, ui_target="ODRZUCONY — brak jawnej czynności", value="DO POTWIERDZENIA", expected_result="DO POTWIERDZENIA", condition_warning="Nie znaleziono jawnej akcji; nie traktować jako instrukcji.", time_range=(start, end), evidence=[], qualification=result["classification"], status="REJECTED", source_segment_ids=candidate.source_segment_ids, frame_ids=candidate.frame_ids, curator_decision=candidate.curator_decision, enriched_step_id=candidate.enriched_step_id, semantic_intent=candidate.semantic_intent, procedure_id=candidate.procedure_id)
            rejected_candidates.append(rejected)
            steps.append(rejected)  # canonical accounting; renderer excludes it
            continue
        candidates_for_step = evidence_rows[index] if index < len(evidence_rows) else []
        candidates_for_step = [
            c for c in candidates_for_step
            if isinstance(c, dict)
            and isinstance(c.get("timestamp"), (int, float))
            and start <= c["timestamp"] <= end
            # A timestamp alone is not provenance.  The selected frame must be
            # explicitly identified and belong to the enriched candidate.
            and c.get("frame_id", c.get("frame_identifier")) in candidate.frame_ids
        ]
        # Evidence is valid only when the enriched candidate maps it to a frame.
        selected = select_evidence_frame(candidates_for_step) if candidate.frame_ids else None
        evidence = []
        if selected:
            evidence.append(Evidence(
                frame_path=str(selected.get("frame_path", selected.get("frame", ""))),
                timestamp=float(selected.get("timestamp", 0)), frame_qa="ok",
                semantic_match=selected.get("semantic_relation") is True,
                ocr_match=selected.get("ocr_relation") is True,
                qualification=result["classification"],
            ))
        qualification = result["classification"]
        status = "CONFIRMED" if qualification == CONFIRMED_ACTION and evidence else "PARTIAL"
        steps.append(ManualStepContract(
            action=text, ui_target=_ui_target(text), value="DO POTWIERDZENIA",
            expected_result="DO POTWIERDZENIA", condition_warning=(None if evidence else "Brak zaakceptowanego dowodu; rezultat wymaga weryfikacji."),
            time_range=(start, end), evidence=evidence, qualification=qualification, status=status,
            source_segment_ids=candidate.source_segment_ids, frame_ids=candidate.frame_ids,
            curator_decision=candidate.curator_decision, enriched_step_id=candidate.enriched_step_id,
            evidence_status="approved" if evidence else "missing", semantic_intent=candidate.semantic_intent,
            procedure_id=candidate.procedure_id,
        ))
    groups: List[ProcedureGroup] = []
    for step in steps:
        key = step.procedure_id or step.semantic_intent
        if key is None:
            groups.append(ProcedureGroup(title="Procedura do potwierdzenia", steps=[step]))
        elif groups and groups[-1].title == key:
            groups[-1].steps.append(step)
        else:
            groups.append(ProcedureGroup(title=key, steps=[step]))
    completeness = "COMPLETE" if steps and all(s.status == "CONFIRMED" and s.evidence and s.expected_result != "DO POTWIERDZENIA" for s in steps) else ("PARTIAL" if steps else "UNKNOWN")
    return ManualDocument(title=title, steps=steps, procedure_groups=groups, goal=goal,
                          preconditions=preconditions or ["DO POTWIERDZENIA"], completeness=completeness,
                          status="DRAFT", rejected_candidates=rejected_candidates, source_candidate_count=len(candidates))


def _client_steps(document: ManualDocument):
    return [step for step in document.all_steps() if step.status in {"CONFIRMED", "PARTIAL"}]


def _safe_frame_ref(frame_path: str) -> str:
    """Expose only a portable filename, never a host filesystem path."""
    return Path(str(frame_path)).name


def render_manual_v2_markdown(document: ManualDocument, metadata: Optional[dict] = None, reference_frames: Optional[Iterable[str]] = None) -> str:
    reference_frames = list(reference_frames or [])
    lines = [f"# {document.title}"]
    accounting = document.accounting()
    lines += ["", f"**Rozliczenie:** źródła {accounting['source_candidates']} · kroki klient-ready {accounting['client_steps']} · odrzucone wewnętrznie {accounting['rejected_internal']}"]
    if metadata:
        labels = {
            "client": "Klient", "process": "Proces", "module": "Moduł",
            "environment": "Środowisko", "author": "Autor",
        }
        visible = [f"**{labels.get(key, key)}:** {value}" for key, value in metadata.items()
                   if value and key in labels]
        if visible:
            lines += ["", "> " + " | ".join(visible)]
    lines += ["", f"**Status:** DRAFT — **Kompletność:** {document.completeness}", "",
             f"## Cel\n{document.goal}", "", "## Warunki wstępne"]
    lines.extend(f"- {item}" for item in document.preconditions)
    lines += ["", "## Procedura"]
    if not document.procedure_groups and document.steps:
        document.procedure_groups = [ProcedureGroup("Procedura", document.steps)]
    number = 0
    rejected = []
    operational_total = sum(step.status in {"CONFIRMED", "PARTIAL"} for step in document.steps)
    if not operational_total:
        lines += ["", "> **DO POTWIERDZENIA:** nie znaleziono wystarczająco wiarygodnej czynności ERP/UI. Odrzucone fragmenty pozostają wyłącznie w raporcie wewnętrznym."]
    for group in document.procedure_groups or [ProcedureGroup("Procedura", document.steps)]:
        operational = [step for step in group.steps if step.status in {"CONFIRMED", "PARTIAL"}]
        if operational:
            lines += ["", f"### {group.title}"]
        for step in operational:
            number += 1
            # Stable compatibility anchor for integrations; numbering remains canonical.
            lines.append(f"### Krok {number}")
            evidence = step.evidence[0] if step.evidence else None
            lines += [f"{number}. **{step.action}** — {step.ui_target}",
                      f"   - Wartość: {step.value}", f"   - Oczekiwany rezultat: {step.expected_result}",
                      f"   - Status: **{step.status}** / kwalifikacja: `{step.qualification}`",
                      f"   - Czas: {step.time_range[0]:.1f}–{step.time_range[1]:.1f}s"]
            if evidence:
                lines.append(f"   - Dowód: `{_safe_frame_ref(evidence.frame_path)}` ({evidence.timestamp:.1f}s)")
            else:
                lines.append("   - Dowód: **DO POTWIERDZENIA** — brak zaakceptowanej klatki")
            if step.condition_warning: lines.append(f"   - Uwaga: {step.condition_warning}")
        rejected.extend(step for step in group.steps if step.status == "REJECTED")
    # REJECTED candidates are internal-only and never leak into client-ready output.
    # Accounting is emitted as metadata, not as rejected-observation content.
    # Client-ready markdown must not expose rejected text or an audit section.
    lines += ["", "## Rezultat końcowy", "DO POTWIERDZENIA — nie potwierdzono kompletnego rezultatu końcowego."]
    return "\n".join(lines) + "\n"


def render_manual_v2_html(document: ManualDocument) -> str:
    import html
    accounting = document.accounting()
    parts = [f"<!doctype html><html lang='pl'><meta charset='utf-8'><title>{html.escape(document.title)}</title><main><h1>{html.escape(document.title)}</h1><p>Status: DRAFT — Kompletność: {html.escape(document.completeness)}</p><p>Rozliczenie: źródła {accounting['source_candidates']} · kroki klient-ready {accounting['client_steps']} · odrzucone wewnętrznie {accounting['rejected_internal']}</p>"]
    n = 0
    for step in document.steps:
        if step.status not in {'CONFIRMED', 'PARTIAL'}: continue
        n += 1; group_title = next((g.title for g in document.procedure_groups if step in g.steps), "Procedura"); parts.append(f"<section data-step-id='{html.escape(step.enriched_step_id or str(n))}' data-status='{step.status}'><h2>{html.escape(group_title)}</h2><h3>Krok {n}</h3><p><b>{html.escape(step.action)}</b> — {html.escape(step.ui_target)}</p><p>Status: {step.status}; Czas: {step.time_range[0]:.1f}–{step.time_range[1]:.1f}s</p>")
        if step.evidence:
            ev = step.evidence[0]; path = html.escape(_safe_frame_ref(ev.frame_path), quote=True); parts.append(f"<img src='frames/{path}' alt='Dowód kroku {n}'>")
        else: parts.append("<p>Dowód: DO POTWIERDZENIA — brak zaakceptowanej klatki</p>")
        parts.append("</section>")
    parts.append("</main></html>"); return "".join(parts)


def render_manual_v2_docx(document: ManualDocument, output_path: Path) -> None:
    """Render the same filtered contract as MD/HTML; no rejected or fallback images."""
    try:
        from docx import Document
        from docx.shared import Inches
    except ImportError:
        return
    accounting = document.accounting()
    doc = Document(); doc.add_heading(document.title, 0); doc.add_paragraph(f"Status: DRAFT — Kompletność: {document.completeness}"); doc.add_paragraph(f"Rozliczenie: źródła {accounting['source_candidates']} · kroki klient-ready {accounting['client_steps']} · odrzucone wewnętrznie {accounting['rejected_internal']}")
    n = 0
    for group in document.procedure_groups or [ProcedureGroup("Procedura", document.steps)]:
        operational = [step for step in group.steps if step.status in {'CONFIRMED', 'PARTIAL'}]
        if not operational:
            continue
        doc.add_heading(str(group.title), level=1)
        for step in operational:
            n += 1
            doc.add_heading(f"Krok {n}", level=2)
            doc.add_paragraph(f"{step.action} — {step.ui_target}")
            doc.add_paragraph(f"Wartość: {step.value}")
            doc.add_paragraph(f"Oczekiwany rezultat: {step.expected_result}")
            doc.add_paragraph(f"Status: {step.status}; ID: {step.enriched_step_id or n}")
            doc.add_paragraph(f"Czas: {step.time_range[0]:.1f}–{step.time_range[1]:.1f}s")
            if step.evidence:
                p = Path(step.evidence[0].frame_path)
                if p.exists():
                    doc.add_paragraph(f"Dowód: {p.name} ({step.evidence[0].timestamp:.1f}s)")
                    doc.add_picture(str(p), width=Inches(6))
            else:
                doc.add_paragraph("Dowód: DO POTWIERDZENIA — brak zaakceptowanej klatki")
    doc.add_heading("Rezultat końcowy", level=1)
    doc.add_paragraph("DO POTWIERDZENIA — nie potwierdzono kompletnego rezultatu końcowego.")
    doc.save(str(output_path))
